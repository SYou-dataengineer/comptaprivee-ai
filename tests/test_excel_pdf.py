"""Excel COM contract tests; no Office installation required in CI."""

import sys
from types import ModuleType, SimpleNamespace
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook
import pytest

from src.comptaprivee import document_converter as converter


class Cells:
    """Sparse Find with Excel's exclusive After and wraparound semantics."""

    def __init__(self, values):
        self.values = values

    def __call__(self, row, column):
        return SimpleNamespace(Row=row, Column=column)

    def Find(self, *, After, SearchOrder, SearchDirection, **kwargs):
        def key(position):
            row, column = position
            return (row, column) if SearchOrder == 1 else (column, row)

        positions = sorted(self.values, key=key)
        if not positions:
            return None
        after = key((After.Row, After.Column))
        if SearchDirection == 1:
            candidates = [p for p in positions if key(p) > after]
            position = candidates[0] if candidates else positions[0]
        else:
            candidates = [p for p in positions if key(p) < after]
            position = candidates[-1] if candidates else positions[-1]
        return self(*position)


def install_com(monkeypatch, values, print_area="", export_error=None):
    sheet = SimpleNamespace(
        Visible=-1, Cells=Cells(values), Rows=SimpleNamespace(Count=1048576),
        Columns=SimpleNamespace(Count=16384),
        PageSetup=SimpleNamespace(PrintArea=print_area),
    )
    ranges = []

    def make_range(first, last):
        bounds = (first.Row, first.Column, last.Row, last.Column)
        zone = SimpleNamespace(
            Address=str(bounds), Width=140, Height=30,
            Columns=SimpleNamespace(AutoFit=Mock()), bounds=bounds,
        )
        ranges.append(zone)
        return zone

    sheet.Range = make_range
    book = Mock(Worksheets=[sheet])
    book.ExportAsFixedFormat.side_effect = export_error
    app = Mock()
    app.Workbooks.Open.return_value = book
    app.InchesToPoints.side_effect = lambda inches: inches * 72
    package = ModuleType("win32com")
    client = ModuleType("win32com.client")
    package.client = client
    client.DispatchEx = Mock(return_value=app)
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setattr(converter, "sys", SimpleNamespace(platform="win32"))
    return app, book, sheet, ranges, client


@pytest.mark.parametrize("values,expected", [
    ({(1, 1): "TEST EXCEL COMPTAPRIVEE"}, (1, 1, 1, 1)),
    ({(2, 2): 123.45}, (2, 2, 2, 2)),
    ({(1, 1): "TEST EXCEL COMPTAPRIVEE", (2, 2): 123.45}, (1, 1, 2, 2)),
    ({(1, 1): "Titre", (2, 2): 123.45, (4, 3): "Autre"}, (1, 1, 4, 3)),
    ({(2, 4): "Titre", (8, 2): 123.45}, (2, 2, 8, 4)),
])
@pytest.mark.parametrize("print_area", ["", "$B$2"])
def test_excel_pdf_includes_all_content(tmp_path, monkeypatch, values, expected, print_area):
    source = tmp_path / "fictif.xlsx"
    wb = Workbook()
    for (row, column), value in values.items():
        wb.active.cell(row, column, value)
    if print_area:
        wb.active.print_area = print_area
    wb.save(source)
    wb.close()
    original = source.read_bytes()
    app, book, sheet, ranges, client = install_com(monkeypatch, values, print_area)
    destination = tmp_path / "fictif.pdf"

    result = converter.convertir_document("Excel → PDF", source, destination)

    assert result.destination == destination
    assert len(ranges) == 1
    assert ranges[0].bounds == expected
    assert sheet.PageSetup.PrintArea == str(expected)
    ranges[0].Columns.AutoFit.assert_called_once_with()
    assert sheet.PageSetup.Zoom is False
    assert sheet.PageSetup.FitToPagesWide == sheet.PageSetup.FitToPagesTall == 1
    client.DispatchEx.assert_called_once_with("Excel.Application")
    app.Workbooks.Open.assert_called_once_with(str(source.resolve()), ReadOnly=True, UpdateLinks=0)
    book.ExportAsFixedFormat.assert_called_once_with(0, str(destination.resolve()), 0, True, False)
    book.Close.assert_called_once_with(False)
    app.Quit.assert_called_once_with()
    assert source.read_bytes() == original
    reloaded = load_workbook(source)
    assert len(reloaded.worksheets) == 1
    reloaded.close()


def test_excel_pdf_office_absent(tmp_path, monkeypatch):
    source = tmp_path / "fictif.xlsx"
    source.touch()
    app, book, sheet, ranges, client = install_com(monkeypatch, {})
    client.DispatchEx.side_effect = RuntimeError("Class not registered")
    with pytest.raises(converter.ErreurConversion, match="Microsoft Excel installé et accessible"):
        converter.excel_vers_pdf(source, tmp_path / "fictif.pdf")
    app.Workbooks.Open.assert_not_called()


def test_excel_pdf_closes_after_export_failure(tmp_path, monkeypatch):
    source = tmp_path / "fictif.xlsx"
    source.touch()
    app, book, *_ = install_com(monkeypatch, {(1, 1): "Titre"}, export_error=RuntimeError("Document not saved"))
    with pytest.raises(converter.ErreurConversion, match="Impossible d'enregistrer le PDF"):
        converter.excel_vers_pdf(source, tmp_path / "fictif.pdf")
    book.Close.assert_called_once_with(False)
    app.Quit.assert_called_once_with()


def test_excel_pdf_empty_sheet_closes(tmp_path, monkeypatch):
    source = tmp_path / "fictif.xlsx"
    source.touch()
    app, book, *_ = install_com(monkeypatch, {})
    with pytest.raises(converter.ErreurConversion, match="Aucune feuille Excel visible"):
        converter.excel_vers_pdf(source, tmp_path / "fictif.pdf")
    book.ExportAsFixedFormat.assert_not_called()
    book.Close.assert_called_once_with(False)
    app.Quit.assert_called_once_with()
